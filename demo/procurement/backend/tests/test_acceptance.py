import asyncio
import json

from docx import Document

from app.evaluate import evaluate
from app.models import BidExtract, ItemExtract, ProjectExtract
from app.parsers import parse_file


def test_reference_pointers_are_not_brands_or_models():
    item = ItemExtract(brand='详见附件3供货范围明细表', model='按实际型号', category='/')
    assert item.brand == '' and item.category == ''
    assert ItemExtract(brand='见投标文件', model='详见附件').model == ''


def test_model_null_placeholders_preserve_missing_semantics():
    record = ProjectExtract.model_validate({'project_name': None, 'warnings': None,
        'lots': [{'lot_name': '包1', 'bids': [{'bidder': '甲', 'evaluation_scores': None, 'source_file': None}],
                  'items': [{'product_name': '设备', 'brand': None, 'quantity': None}]}]})
    assert record.project_name == '' and record.warnings == []
    assert record.lots[0].bids[0].evaluation_scores == []
    assert record.lots[0].items[0].brand == '' and record.lots[0].items[0].quantity is None


def test_disclosed_scores_support_participation_but_never_override_failure():
    bid = BidExtract(bidder='甲', evaluation_scores=[81, 83, 82])
    assert bid.eligible and bid.review_status == '未披露'
    failed = BidExtract(bidder='乙', evaluation_scores=[0], qualification_status='未通过')
    assert not failed.eligible


def test_real_provider_labelled_score_records_preserve_components():
    bid = BidExtract(bidder='甲', evaluation_scores=[
        {'item': '技术得分', 'value': 39.67},
        {'item': '商务得分', 'value': '27.0'},
        {'item': '价格得分', 'value': 23.41},
    ])
    assert bid.evaluation_scores[0] == {'item': '技术得分', 'value': 39.67}
    assert bid.evaluation_scores[1]['value'] == 27.0
    assert bid.eligible and bid.score is None
    assert json.loads(bid.model_dump_json())['evaluation_scores'][2]['value'] == 23.41
    unknown = BidExtract(bidder='乙', evaluation_scores=[None, {'item': '技术得分', 'value': None}])
    assert unknown.evaluation_scores == [] and not unknown.eligible


def test_docx_source_rows_distinguish_tables(tmp_path):
    path = tmp_path / 'two-tables.docx'
    doc = Document()
    for name in ('商品甲', '商品乙'):
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = name
    doc.save(path)
    text = parse_file(path)
    assert '[TABLE:1 ROW:1]' in text and '[TABLE:2 ROW:1]' in text


def test_evaluation_uses_uuid_not_file_order(tmp_path):
    gold = [{'notice_uuid': 'n1', 'items': [{'product_name': '甲'}]},
            {'notice_uuid': 'n2', 'items': [{'product_name': '乙'}]}]
    gp, pp = tmp_path/'reference.jsonl', tmp_path/'prediction.jsonl'
    gp.write_text('\n'.join(json.dumps(r) for r in gold), encoding='utf-8')
    pp.write_text('\n'.join(json.dumps(r) for r in reversed(gold)), encoding='utf-8')
    result = asyncio.run(evaluate(gp, pp))
    assert result['micro']['f1'] == 1


def test_missing_predictions_count_as_missing_rows(tmp_path):
    gold = [{'notice_uuid': 'n1', 'items': [{'product_name': '甲', 'quantity': 2}]}]
    gp, pp = tmp_path/'reference.jsonl', tmp_path/'prediction.jsonl'
    gp.write_text(json.dumps(gold[0]), encoding='utf-8')
    pp.write_text('', encoding='utf-8')
    result = asyncio.run(evaluate(gp, pp))
    assert result['missing_notice_ids'] == ['n1']
    assert result['micro']['recall'] == 0 and result['diagnostics'][0]['missing_rows'] == 1
