import asyncio
from dataclasses import replace
import io
import json
import math
from pathlib import Path
import sqlite3
import zipfile

import httpx
from openpyxl import load_workbook
import pytest
from docx import Document
from pypdf import PdfWriter

from app import analytics,database,jobs,ingestion,main
from app.config import settings
from app.database import connect,init_db,one,rows,upsert_notice
from app.evaluate import evaluate,score_items
from app.exports import HEADERS,workbook_bytes,export_records
from app.extractor import DeepSeekExtractor,chunk_text,merge_results,validate_result
from app.models import BidExtract,ItemExtract,LotExtract,ProjectExtract
from app.normalization import money_to_yuan,normalize_org,yuan_to_wan
from app.parsers import ArchiveBudget,ParseError,detect_type,expand_archives,group_documents,parse_file,safe_unzip

@pytest.mark.parametrize('raw,expected',[('369.5015(万元)',3695015),('1.2万',12000),('0.3亿',30000000),('￥1,234元',1234),(0,0),(float('nan'),None)])
def test_amount_units(raw,expected):
    assert money_to_yuan(raw)==expected

def test_half_up_and_item_precision():
    assert yuan_to_wan(945350,2)==94.54
    assert yuan_to_wan(930)==.093

def test_organization_parentheses():
    assert normalize_org('某 公司（广东）')==normalize_org('某公司(广东)')

@pytest.mark.parametrize('state,qual,comp,winner,eligible',[('通过','未披露','未披露',False,True),('未披露','通过','通过',False,True),('通过','未通过','通过',True,False),('未披露','未披露','未披露',False,False),('未披露','未披露','未披露',True,True)])
def test_audit_eligibility(state,qual,comp,winner,eligible):
    bid=BidExtract(bidder='甲',review_status=state,qualification_status=qual,compliance_status=comp,is_winner=winner)
    assert bid.eligible is eligible
    if not eligible and qual=='未通过': assert bid.role=='其他'

def test_nan_confidence_does_not_become_certainty():
    assert ItemExtract(confidence=float('nan')).confidence==0

def test_idempotent_upsert_and_namespace(notice):
    record=notice()
    first=upsert_notice(record,['a.html'])
    second=upsert_notice(record,['a.html'])
    assert first==second
    assert one('SELECT COUNT(*) n FROM projects')['n']==1
    assert one('SELECT COUNT(*) n FROM items')['n']==1
    upsert_notice(record,['a.html'],'official')
    assert one('SELECT COUNT(*) n FROM projects')['n']==2

def test_cross_lot_does_not_form_common_bid():
    result=ProjectExtract(notice_uuid='n',project_name='项目',purchaser='买方',lots=[
        LotExtract(lot_name='包1',awarded_amount=100,bids=[BidExtract(bidder='甲',is_winner=True)]),
        LotExtract(lot_name='包2',awarded_amount=200,bids=[BidExtract(bidder='乙',is_winner=True)])])
    upsert_notice(result,['n.html'])
    assert analytics.common_projects(['甲','乙'])['summary']['project_count']==0
    assert analytics.common_purchasers(['甲','乙'])['purchasers'][0]['total_amount']==300

def test_explicit_failed_bid_kept_but_excluded(notice):
    result=notice(); result.bids.append(BidExtract(bidder='失败企业',qualification_status='未通过',bid_amount=1))
    upsert_notice(result,['n.html'])
    assert one("SELECT COUNT(*) n FROM bids WHERE bidder='失败企业'")['n']==1
    assert analytics.common_projects(['供应商甲','失败企业'])['summary']['project_count']==0
    assert '失败企业' not in {n['name'] for n in analytics.relationship_graph()['nodes']}

def test_unknown_lot_not_assumed_same_package(notice):
    upsert_notice(notice(lot=''),['n.html'])
    assert analytics.common_projects(['供应商甲','供应商乙'])['summary']['project_count']==0

def test_fallback_not_in_relationships(notice):
    upsert_notice(notice(processing_status='fallback'),['n.html'])
    assert analytics.relationship_graph()['nodes']==[]

def test_brand_cooperation_is_not_item_count(notice):
    result=notice(); result.items.append(ItemExtract(product_name='另一项',brand='品牌X',total_price=5))
    upsert_notice(result,['n.html'])
    brand=analytics.purchaser_partners('采购单位A')['product_suppliers'][0]
    assert brand['cooperation_count']==1 and brand['item_count']==2

def test_missing_amount_not_zero(notice):
    upsert_notice(notice(amount=None),['n.html'])
    assert analytics.common_projects(['供应商甲','供应商乙'])['summary']['total_amount'] is None

def test_consortium_total_not_credited_to_each_member():
    record=ProjectExtract(notice_uuid='n',project_name='联合体',purchaser='买方',lot_name='包1',awarded_amount=300,bids=[
        BidExtract(bidder='甲',is_winner=True,bid_amount=300,consortium_id='联合体1'),
        BidExtract(bidder='乙',is_winner=True,bid_amount=300,consortium_id='联合体1')])
    upsert_notice(record,['n.html'])
    result=analytics.common_purchasers(['甲','乙'])['purchasers'][0]
    assert result['total_amount']==300
    assert all(x['total_amount'] is None for x in result['supplier_details'])

def test_superseded_lot_removed(notice):
    record=notice(); record.lots=[LotExtract(lot_name='包1'),LotExtract(lot_name='包2')]
    upsert_notice(record,['n.html'])
    record.lots=record.lots[:1]; upsert_notice(record,['n.html'])
    assert one('SELECT COUNT(*) n FROM projects')['n']==1

def test_migration_keeps_legacy_and_creates_backup(tmp_path,monkeypatch):
    target=tmp_path/'legacy.db'
    with sqlite3.connect(target) as c:
        c.execute("CREATE TABLE projects(id INTEGER PRIMARY KEY AUTOINCREMENT,project_code TEXT NOT NULL DEFAULT '',name TEXT NOT NULL,purchaser TEXT NOT NULL DEFAULT '',budget_amount REAL,awarded_amount REAL,announcement_date TEXT NOT NULL DEFAULT '',source_files TEXT NOT NULL DEFAULT '',extraction_method TEXT NOT NULL DEFAULT '',warnings TEXT NOT NULL DEFAULT '',created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)")
        c.execute("INSERT INTO projects(name,source_files) VALUES('原项目','notice.html')")
    monkeypatch.setattr(database,'settings',replace(settings,database_path=target))
    init_db(); init_db()
    assert one('SELECT COUNT(*) n FROM projects')['n']==1
    assert one('SELECT notice_uuid FROM projects')['notice_uuid']=='notice'
    assert target.with_suffix('.pre-v2.db').is_file()

def test_raw_missing_prices_preserved(notice):
    result=notice(); result.items=[ItemExtract(product_name='服务',quantity=2,total_price=100)]
    validate_result(result)
    assert result.items[0].unit_price is None
    assert result.items[0].derived_values['unit_price']==50

def test_distinct_supplier_rows_not_merged():
    records=[ProjectExtract(lot_name='包1',items=[ItemExtract(product_name='相同名称',model='M',supplier=s,source_file='doc',line_id='1')]) for s in ('甲','乙')]
    assert len(merge_results(records).items)==2

def test_repeated_source_row_merges_missing_fields():
    records=[ProjectExtract(lot_name='包1',items=[ItemExtract(product_name='产品',source_file='doc',line_id='1',brand=b)]) for b in ('','品牌')]
    result=merge_results(records)
    assert len(result.items)==1 and result.items[0].brand=='品牌'

def test_conflicting_audit_failure_dominates():
    result=merge_results([ProjectExtract(lot_name='包1',bids=[BidExtract(bidder='甲',review_status=s)]) for s in ('通过','未通过')])
    assert not result.bids[0].eligible

def test_chunking_keeps_all_characters():
    text=('一段完整的话。\n'*1000)
    parts=chunk_text(text,500)
    assert ''.join(parts)==text and max(map(len,parts))<=500

@pytest.mark.parametrize('member',['../escape.txt','/absolute.txt','C:/escape.txt','folder/../../escape.txt'])
def test_zip_path_traversal_rejected(tmp_path,member):
    archive=tmp_path/'input.zip'
    with zipfile.ZipFile(archive,'w') as z: z.writestr(member,'data')
    with pytest.raises(ParseError): safe_unzip(archive,tmp_path/'out')

def test_archive_budget_shared_across_nested_files(tmp_path):
    inner=io.BytesIO()
    with zipfile.ZipFile(inner,'w') as z: z.writestr('file.txt','a'*100)
    archive=tmp_path/'outer.zip'
    with zipfile.ZipFile(archive,'w') as z: z.writestr('inner.zip',inner.getvalue())
    budget=ArchiveBudget(limit=len(inner.getvalue())+90)
    child=safe_unzip(archive,tmp_path/'out',budget=budget)[0]
    with pytest.raises(ParseError): safe_unzip(child,tmp_path/'inner',budget=budget)

def test_docx_disguised_as_doc_detected_and_parsed(tmp_path):
    path=tmp_path/'old.doc'; d=Document(); d.add_paragraph('真实文档内容'); d.save(path)
    assert detect_type(path)=='.docx'
    assert '真实文档内容' in parse_file(path)

def test_maintenance_response_is_not_document(tmp_path):
    path=tmp_path/'bad.doc'; path.write_text('<html><body>系统正在维护中</body></html>',encoding='utf-8')
    with pytest.raises(ParseError): parse_file(path)

def test_word_html_table_is_usable(tmp_path):
    path=tmp_path/'document.doc'; path.write_text('<html><table><tr><td>真实项目</td><td>123</td></tr></table></html>',encoding='utf-8')
    assert '真实项目 | 123' in parse_file(path)

def test_pdf_disguised_as_docx_is_detected(tmp_path):
    path=tmp_path/'document.docx'; w=PdfWriter(); w.add_blank_page(width=100,height=100); w.write(str(path))
    assert detect_type(path)=='.pdf'
    with pytest.raises(ParseError): parse_file(path)

def test_html_line_break_inside_model_does_not_invent_a_space(tmp_path):
    path = tmp_path / 'notice.html'
    path.write_text('<table><tr><td>终端管理器</td><td><span>DH-ITSE0400-TA</span><br> <span>-N02</span></td><td>Model A B</td></tr></table>', encoding='utf-8')
    text = parse_file(path)
    assert '终端管理器 | DH-ITSE0400-TA-N02 | Model A B' in text


def test_html_merged_cells_keep_row_ownership(tmp_path):
    path=tmp_path/'test.html'; path.write_text('<table><tr><td rowspan="2">包1</td><td>产品甲</td></tr><tr><td>产品乙</td></tr></table>',encoding='utf-8')
    text=parse_file(path)
    assert '包1 | 产品乙' in text

def test_ambiguous_attachment_not_guessed(tmp_path):
    files=[tmp_path/'a.html',tmp_path/'b.html',tmp_path/'other.pdf']
    with pytest.raises(ParseError): group_documents(files)

def test_notice_zip_lazy_and_exact_group(tmp_path):
    html=tmp_path/'t20260101_10001.html'; html.write_text('<html>正文</html>',encoding='utf-8')
    archive=tmp_path/'t20260101_10001.zip'
    with zipfile.ZipFile(archive,'w') as z: z.writestr('large.pdf',b'not read at discovery')
    issues=[]; groups=ingestion.discover_uploads([html,archive],tmp_path/'discover',issues)
    assert groups==[[html,archive]] and not (tmp_path/'discover').exists() and not issues

def test_model_json_path_and_provenance(tmp_path):
    html=tmp_path/'n.html'; html.write_text('<html>项目公告及采购包</html>',encoding='utf-8')
    extractor=DeepSeekExtractor(); extractor.api_key='test-only'
    response={'project_name':'项目','lots':[{'lot_name':'包1','awarded_amount':'2万元','items':[{'product_name':'产品','unit_price':930,'quantity':2,'source_file':'invented'}],'bids':[{'bidder':'甲','is_winner':True}]}]}
    def handler(request):
        assert '包' in json.loads(request.content)['messages'][0]['content']
        return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps(response)}}]})
    async def run():
        extractor.client=httpx.AsyncClient(transport=httpx.MockTransport(handler))
        try: return await extractor.extract([html])
        finally: await extractor.close()
    result=asyncio.run(run())
    assert result.processing_status=='complete' and extractor.calls==1
    assert result.lots[0].awarded_amount==20000
    assert result.lots[0].items[0].source_file=='doc-1:n.html'
    assert result.lots[0].items[0].source=='announcement'

def test_model_invalid_json_retries(monkeypatch):
    async def no_wait(_): pass
    monkeypatch.setattr('app.extractor.asyncio.sleep',no_wait)
    extractor=DeepSeekExtractor(); extractor.api_key='test-only'; count=0
    def handler(request):
        nonlocal count
        count+=1
        return httpx.Response(200,json={'choices':[{'message':{'content':'bad' if count==1 else '{"project_name":"项目"}'}}]})
    async def run():
        extractor.client=httpx.AsyncClient(transport=httpx.MockTransport(handler))
        try: return await extractor._call('正文',['doc'])
        finally: await extractor.close()
    assert asyncio.run(run()).project_name=='项目'
    assert count==2

def test_three_sheet_columns_units_and_missing(notice):
    record=notice(amount=945350); record.items[0].brand=''; record.bids[1].qualification_status='未通过'; record.bids[1].review_reason='资格不足'
    w=load_workbook(io.BytesIO(workbook_bytes([record])))
    for name,headers in HEADERS.items(): assert [c.value for c in w[name][1]]==headers
    assert w['总表']['H2'].value==94.54
    assert w['标的表']['I2'].value==.093
    assert w['标的表']['F2'].value is None
    assert w['标的表']['B2'].value==record.notice_uuid
    assert w['投标人表']['E3'].value=='其他'
    assert '资格不足' in w['投标人表']['F3'].value

def test_formula_injection_not_executable(notice):
    record=notice(); record.project_name='=HYPERLINK("evil")'
    w=load_workbook(io.BytesIO(workbook_bytes([record])))
    assert w['总表']['E2'].data_type=='s'

def test_job_persistence_and_snapshot_is_stable(notice,tmp_path):
    record=notice(); state=jobs.create_job('job1',[],tmp_path,groups=[])
    jobs._record(state.id,record.notice_uuid,'complete',record.model_dump(),[],1.2)
    with database.connect() as c: c.execute("UPDATE jobs SET status='completed' WHERE id='job1'")
    record.project_name='修改后的项目'; upsert_notice(record,['n.html'])
    snapshots,_=export_records('default','job1')
    assert snapshots[0].project_name=='采购项目'
    assert jobs.get_job('job1').finished

def test_restart_marks_interrupted(tmp_path):
    jobs.create_job('job1',[],tmp_path,groups=[]); jobs.interrupt_running_jobs()
    assert jobs.get_job('job1').status=='interrupted'

def test_active_batch_exclusion(tmp_path):
    jobs.create_job('one',[],tmp_path,groups=[])
    with pytest.raises(ValueError): jobs.create_job('two',[],tmp_path,groups=[])

def test_parse_only_does_not_call_model_or_write_business(tmp_path,monkeypatch):
    html=tmp_path/'t20260101_10001.html'; html.write_text('<html><body>项目名称：测试项目</body></html>',encoding='utf-8')
    async def forbidden(*args): raise AssertionError('不允许预检调用模型')
    monkeypatch.setattr(DeepSeekExtractor,'extract',forbidden)
    jobs.create_job('parse',[],tmp_path/'job','official',[[html]],'parse-only')
    asyncio.run(jobs.process_job('parse'))
    assert jobs.get_job('parse').status=='completed'
    assert one('SELECT COUNT(*) n FROM projects')['n']==0
    with pytest.raises(ValueError): export_records('official','parse')

def test_no_key_is_partial_fallback_not_success(tmp_path):
    html=tmp_path/'t20260101_10001.html'; html.write_text('<html><body>项目名称：测试项目\n采购单位：示例大学</body></html>',encoding='utf-8')
    jobs.create_job('fallback',[],tmp_path/'job','official',[[html]])
    asyncio.run(jobs.process_job('fallback'))
    state=jobs.get_job('fallback')
    assert state.status=='completed_with_warnings' and state.succeeded==0 and state.partial==1
    assert one("SELECT COUNT(*) n FROM projects WHERE dataset_id='official'")['n']==1

def test_unchanged_complete_inputs_cached(tmp_path,monkeypatch,notice):
    html=tmp_path/'t20260101_10001.html'; html.write_text('<html>项目完整正文</html>',encoding='utf-8')
    calls=0
    async def extract(self,files):
        nonlocal calls
        calls+=1; return notice()
    monkeypatch.setattr(DeepSeekExtractor,'extract',extract)
    for jobid in ('first','second'):
        jobs.create_job(jobid,[],tmp_path/jobid,'official',[[html]])
        asyncio.run(jobs.process_job(jobid))
    assert calls==1 and jobs.get_job('second').skipped==1

def test_swapped_price_not_perfect_score():
    gold=[{'product_name':'甲产品','unit_price':100},{'product_name':'乙产品','unit_price':200}]
    predicted=[{'product_name':'甲产品','unit_price':200},{'product_name':'乙产品','unit_price':100}]
    report=score_items(gold,predicted)
    assert report['fields']['unit_price']['tp']==0
    assert report['fields']['unit_price']['wrong_value']==2
    assert report['micro']['precision']<1

def test_duplicate_rows_are_not_sets():
    record={'product_name':'产品','unit_price':100}
    report=score_items([record,record],[record])
    assert report['missing_rows']==1 and report['fields']['unit_price']['fn']==1
    assert report['warnings']

def test_missing_true_negative_and_integer_float():
    report=score_items([{'product_name':'产品','unit_price':100}],[{'product_name':'产品','unit_price':100.0}])
    assert report['fields']['unit_price']['tp']==1
    assert report['fields']['brand']['tn']==1

def test_evaluation_requires_explicit_predictions(tmp_path):
    gold=tmp_path/'gold.jsonl'; gold.write_text('{"items":[]}',encoding='utf-8')
    with pytest.raises(ValueError): asyncio.run(evaluate(gold))

@pytest.mark.parametrize('url',['/api/projects?page=0','/api/items?page_size=-1','/api/items?min_price=-1','/api/entities?type=invalid','/api/graph?limit=0','/api/stats?dataset_id=../outside'])
def test_api_input_validation(client,url):
    assert client.get(url).status_code==422

def test_model_filter_ignores_whitespace_and_keeps_source_text(client,notice):
    record = notice()
    record.items[0] = record.items[0].model_copy(update={'model': 'DH-ITSE0400-TA -N02'})
    upsert_notice(record, ['n.html'], 'official')
    stored = one('SELECT model, model_key FROM items')
    assert stored['model'] == 'DH-ITSE0400-TA -N02'
    assert stored['model_key'] == 'DH-ITSE0400-TA-N02'
    exact = client.get('/api/items', params={'dataset_id': 'official', 'model': 'DH-ITSE0400-TA-N02'})
    spaced = client.get('/api/items', params={'dataset_id': 'official', 'model': 'DH-ITSE0400-TA -N02'})
    other = client.get('/api/items', params={'dataset_id': 'official', 'model': 'DH-ITSE0400-TB-N02'})
    assert exact.json()['total'] == spaced.json()['total'] == 1
    assert exact.json()['data'][0]['model'] == 'DH-ITSE0400-TA -N02'
    assert other.json()['total'] == 0


def test_existing_model_rows_get_match_keys(tmp_path, monkeypatch):
    target = tmp_path / 'legacy.db'
    monkeypatch.setattr(database, 'settings', replace(settings, database_path=target))
    init_db()
    with connect() as connection:
        connection.execute("INSERT INTO projects(name) VALUES('旧项目')")
        project_id = connection.execute('SELECT id FROM projects').fetchone()['id']
        connection.execute("INSERT INTO items(project_id, model, model_key) VALUES(?, 'DH-ITSE0400-TA -N02', '')", (project_id,))
    init_db()
    assert one("SELECT model FROM items")['model'] == 'DH-ITSE0400-TA -N02'
    assert one("SELECT model_key FROM items")['model_key'] == 'DH-ITSE0400-TA-N02'


def test_api_numeric_filters_and_dataset_scope(client,notice):
    upsert_notice(notice(),['n.html'],'official')
    assert client.get('/api/items').json()['total']==0
    assert client.get('/api/items?dataset_id=official&min_price=931').json()['total']==0
    assert client.get('/api/items?dataset_id=official&max_price=930&brand=品牌X').json()['total']==1
    assert client.get('/api/items?min_price=10&max_price=1').status_code==400

@pytest.mark.parametrize('query,dataset_id,expected',[
    ('采购单位A','default',1),('供应商甲','default',1),(' 供应商甲 ','default',1),
    ('供应商乙','default',0),('供应商甲','official',0),
])
def test_item_search_matches_purchaser_and_winning_supplier_within_dataset(client,notice,query,dataset_id,expected):
    upsert_notice(notice(),['n.html'])
    response=client.get('/api/items',params={'q':query,'dataset_id':dataset_id})
    assert response.status_code==200
    assert response.json()['total']==expected


def test_auth_covers_reads_writes_and_exports(client,monkeypatch):
    monkeypatch.setattr(main,'settings',replace(settings,api_token='test-token'))
    assert client.get('/api/health').status_code==200
    for endpoint in ('/api/stats','/api/jobs','/api/export/standard.xlsx'):
        assert client.get(endpoint).status_code==401
        assert client.get(endpoint,headers={'Authorization':'Bearer test-token'}).status_code==200
    assert client.post('/api/ingest/local',json={'limit':1,'mode':'parse-only'}).status_code==401

def test_api_export_and_evidence_detail(client,notice):
    ids=upsert_notice(notice(),['n.html'])
    response=client.get('/api/export/standard.xlsx')
    assert response.status_code==200 and response.content[:2]==b'PK'
    assert response.headers['content-disposition']=='attachment; filename=procurement-tables.xlsx'
    detail_export=client.get('/api/export/standard.xlsx',params={'item_view':'detail'})
    assert detail_export.headers['content-disposition']=='attachment; filename=procurement-detail-tables.xlsx'
    detail=client.get(f'/api/projects/{ids[0]}').json()
    assert len(detail['bids'])==2 and len(detail['items'])==1

def test_local_ingest_rejects_arbitrary_path(client):
    assert client.post('/api/ingest/local',json={'notice_ids':['../../outside']}).status_code==400


def test_failed_boolean_audit_is_not_unknown():
    bid=BidExtract(bidder='甲',is_winner=True,review_status=False)
    assert not bid.eligible and bid.role=='其他'

def test_failed_audit_reason_inside_status_preserved():
    bid=BidExtract(bidder='甲',qualification_status='未通过（缺少许可）')
    assert bid.qualification_status=='未通过' and '缺少许可' in bid.review_reason

def test_unknown_purchasers_not_artificially_grouped(notice):
    for uid,name in [('a','甲'),('b','乙')]:
        record=notice(uid=uid); record.purchaser=''; record.bids=[BidExtract(bidder=name,is_winner=True)]
        upsert_notice(record,['n.html'])
    assert analytics.common_purchasers(['甲','乙'])['purchasers']==[]

def test_single_award_amount_takes_precedence_over_initial_quote(notice):
    record=notice(amount=300); record.bids[0].bid_amount=350
    upsert_notice(record,['n.html'])
    assert analytics.purchaser_partners('采购单位A')['winning_suppliers'][0]['total_amount']==300

def test_model_cannot_supply_unverifiable_derived_price(notice):
    record=notice(); record.items=[ItemExtract(product_name='项目',derived_values={'unit_price':999})]
    validate_result(record)
    assert record.items[0].derived_values=={}

def test_concurrent_jobs_only_one_accepted(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    def submit(name):
        try: jobs.create_job(name,[],tmp_path/name,groups=[]); return True
        except ValueError: return False
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(submit,['one','two']))
    assert sum(results)==1

def test_changed_model_endpoint_invalidates_cache(tmp_path,monkeypatch,notice):
    html=tmp_path/'t20260101_10001.html'; html.write_text('<html>项目完整正文</html>',encoding='utf-8')
    calls=0
    async def extract(self,files):
        nonlocal calls
        calls+=1; return notice()
    monkeypatch.setattr(DeepSeekExtractor,'extract',extract)
    jobs.create_job('first',[],tmp_path/'first','official',[[html]])
    asyncio.run(jobs.process_job('first'))
    monkeypatch.setattr(jobs,'settings',replace(settings,deepseek_base_url='https://other.example.test'))
    jobs.create_job('second',[],tmp_path/'second','official',[[html]])
    asyncio.run(jobs.process_job('second'))
    assert calls==2 and jobs.get_job('second').skipped==0

def test_document_size_guard(tmp_path,monkeypatch):
    import app.parsers as parsers
    path=tmp_path/'big.txt'; path.write_text('1234567890',encoding='utf-8')
    monkeypatch.setattr(parsers,'settings',replace(settings,max_document_bytes=5))
    with pytest.raises(ParseError,match='文档'): parse_file(path)

def test_ooxml_internal_size_guard(tmp_path,monkeypatch):
    import app.parsers as parsers
    path=tmp_path/'zipbomb.docx'
    with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('word/document.xml','x'*5000)
    monkeypatch.setattr(parsers,'settings',replace(settings,max_document_bytes=1000))
    with pytest.raises(ParseError,match='内部'): parse_file(path)

def test_uploaded_notice_precheck_runs_via_api(client):
    response=client.post('/api/ingest',data={'dataset_id':'official','mode':'parse-only'},files=[('files',('t20260101_10001.html','<html>项目名称：正确预检</html>','text/html'))])
    assert response.status_code==202
    job=client.get('/api/jobs/'+response.json()['id']).json()
    assert job['mode']=='parse-only' and job['status']=='completed'
    assert client.get('/api/stats?dataset_id=official').json()['projects']==0

def test_upload_limit_rejects_and_cleans(tmp_path,monkeypatch):
    from fastapi import UploadFile
    monkeypatch.setattr(jobs,'settings',replace(settings,max_upload_bytes=2,upload_dir=tmp_path))
    with pytest.raises(ValueError): asyncio.run(jobs.save_uploads([UploadFile(io.BytesIO(b'123'),filename='test.txt')]))
    assert list(tmp_path.iterdir())==[]

def test_duplicate_notice_id_in_batch_fails_without_overwriting_snapshot(tmp_path):
    a,b=tmp_path/'a',tmp_path/'b'; a.mkdir(); b.mkdir()
    for directory in (a,b): (directory/'n.html').write_text('<html>内容</html>',encoding='utf-8')
    jobs.create_job('duplicate',[],tmp_path/'job',groups=[[a/'n.html'],[b/'n.html']],mode='parse-only')
    asyncio.run(jobs.process_job('duplicate'))
    assert jobs.get_job('duplicate').status=='failed'
    assert one('SELECT COUNT(*) n FROM job_documents')['n']==0


def test_strict_export_has_exactly_three_template_sheets(notice):
    record=notice()
    strict=load_workbook(io.BytesIO(workbook_bytes([record])))
    annotated=load_workbook(io.BytesIO(workbook_bytes([record],include_notes=True)))
    assert strict.sheetnames==['总表','标的表','投标人表']
    assert annotated.sheetnames==['总表','标的表','投标人表','处理说明']

def test_separate_quality_export_keeps_warnings(client,notice):
    record=notice(); record.warnings=['需要复核']
    upsert_notice(record,['n.html'])
    response=client.get('/api/export/quality.json')
    assert response.status_code==200
    assert response.json()['notices'][0]['warnings']==['需要复核']


@pytest.mark.parametrize('raw',['1.8%','报价下浮率10%','100-200万元','USD 100','100美元'])
def test_rates_ranges_foreign_currency_not_coerced_to_yuan(raw):
    assert money_to_yuan(raw) is None

@pytest.mark.parametrize('raw,expected',[('.5万元',5000),('1e5元',100000)])
def test_fractional_and_exponent_amounts(raw,expected):
    assert money_to_yuan(raw)==expected


def test_unknown_audit_left_blank_even_for_inferred_eligible_winner(notice):
    record=notice()
    w=load_workbook(io.BytesIO(workbook_bytes([record])))
    assert record.bids[0].eligible
    assert w['投标人表']['E2'].value=='中标人'
    assert w['投标人表']['F2'].value is None

def test_legacy_missing_uuid_not_fabricated():
    with database.connect() as c:
        c.execute("INSERT INTO projects(name,dataset_id) VALUES('历史项目','default')")
    records,_=export_records('default')
    assert records[0].notice_uuid==''
