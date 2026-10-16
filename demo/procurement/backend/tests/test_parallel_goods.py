"""Focused synthetic regressions; no model HTTP or official-cohort reruns."""
import asyncio
import hashlib
import io
import json
from dataclasses import replace
from html import escape

from docx import Document
import httpx
from openpyxl import load_workbook
import pytest

from app.database import rows, upsert_notice
from app.evidence import canonical_location, ground_evidence
from app.evaluate import score_items
from app.exports import export_records, workbook_bytes
from app.extractor import DeepSeekExtractor, FatalModelError, merge_results, validate_result
from app.item_hierarchy import _verified
from app.jobs import create_job, get_job, process_job
from app.models import BidExtract, ItemExtract, LotExtract, ProjectExtract
from app.normalization import money_to_yuan, number_value
from app.parsers import PARSE_WARNING_PREFIX, parse_file
from app.prompts import MULTI_ITEM_RULES, system_prompt, user_prompt


HEADER = ['序号', '供应商名称', '货物名称', '货物品牌', '货物型号', '货物数量', '货物单价(元)']
GOODS = ['1-1', '北京万迪科信息技术有限公司', '光跳纤架；光熔纤柜', '康普；康普',
         'NGFC FRAME；MX6', '4套；5套', '￥319,300.00；￥60,480.00']


def table_file(tmp_path, records, kind='html'):
    path = tmp_path / ('notice.' + kind)
    if kind == 'docx':
        document = Document()
        table = document.add_table(rows=0, cols=len(records[0]))
        for values in records:
            for cell, value in zip(table.add_row().cells, values):
                cell.text = value
        document.save(path)
    else:
        path.write_text('<html><table>' + ''.join(
            '<tr>' + ''.join('<td>' + escape(value) + '</td>' for value in record) + '</tr>'
            for record in records) + '</table></html>', encoding='utf-8')
    return path


@pytest.mark.parametrize('kind', ['html', 'docx'])
@pytest.mark.parametrize('separator', ['；', ';'])
def test_aligned_goods_are_independent_source_rows(tmp_path, kind, separator):
    data = [value.replace('；', separator) for value in GOODS]
    text = parse_file(table_file(tmp_path, [HEADER, data], kind))
    assert '[TABLE:1 ROW:2 ITEM:1] 1-1 | 北京万迪科信息技术有限公司 | 光跳纤架 | 康普 | NGFC FRAME | 4套 | ￥319,300.00' in text
    assert '[TABLE:1 ROW:2 ITEM:2] 1-1 | 北京万迪科信息技术有限公司 | 光熔纤柜 | 康普 | MX6 | 5套 | ￥60,480.00' in text
    assert '[TABLE:1 ROW:2]' not in text
    assert PARSE_WARNING_PREFIX not in text


def test_repeated_headers_update_positions_and_keep_core_labels(tmp_path):
    order = [1, 6, 4, 0, 3, 2, 5]
    data = [*GOODS]
    data[2] = ' 光跳纤架（核心产品） ； 光熔纤柜 '
    text = parse_file(table_file(tmp_path, [HEADER, GOODS, [HEADER[i] for i in order], [data[i] for i in order]]))
    assert '[TABLE:1 ROW:4 ITEM:2] 北京万迪科信息技术有限公司 | ￥60,480.00 | MX6 | 1-1 | 康普 | 光熔纤柜 | 5套' in text
    assert '光跳纤架（核心产品）' in text
    assert text.count(' ITEM:') == 4


@pytest.mark.parametrize('column,value', [
    (0, '1；2'), (1, '供应商甲；供应商乙'), (3, '康普'),
    (4, 'M1；；M2'), (5, '4套；'), (6, '319300；60480；100'),
])
def test_ambiguous_columns_keep_original_and_warn(tmp_path, column, value):
    data = [*GOODS]
    data[column] = value
    text = parse_file(table_file(tmp_path, [HEADER, data]))
    assert '[TABLE:1 ROW:2] ' + ' | '.join(data) in text
    assert ' ITEM:' not in text
    assert PARSE_WARNING_PREFIX in text
    # The warning must not replace the source row in evidence grounding.
    assert canonical_location(text.splitlines()[-1]) == 'TABLE:1 ROW:2'


def test_other_tables_and_incomplete_headers_are_not_split(tmp_path):
    text = parse_file(table_file(tmp_path, [['专家名单', '地址'], ['张三；李四', '地址甲；地址乙']]))
    assert '张三；李四 | 地址甲；地址乙' in text and ' ITEM:' not in text
    header = [*HEADER]
    header[6] = '中标总额'
    text = parse_file(table_file(tmp_path, [header, GOODS]))
    assert '光跳纤架；光熔纤柜' in text and ' ITEM:' not in text


def test_single_goods_and_blank_rows_keep_source_numbering(tmp_path):
    single = [*GOODS]
    single[2:] = ['设备一批', '品牌甲；品牌乙', 'M1', '1套', '100元']
    text = parse_file(table_file(tmp_path, [HEADER, [''] * len(HEADER), single]))
    assert '[TABLE:1 ROW:3] ' + ' | '.join(single) in text
    assert '[TABLE:1 ROW:2]' not in text and ' ITEM:' not in text


@pytest.mark.parametrize('parse,value', [
    (number_value, '4套；5套'), (number_value, '４套;５套'),
    (money_to_yuan, '￥319,300.00；￥60,480.00'),
    (money_to_yuan, '1万元;2万元'),
])
def test_parallel_numbers_never_silently_become_the_first_value(parse, value):
    assert parse(value) is None


@pytest.mark.parametrize('parse,value,expected', [
    (number_value, '4套；', 4), (number_value, '数量：0.5套', .5),
    (money_to_yuan, '￥319,300.00', 319300), (money_to_yuan, '2万元；含税', 20000),
])
def test_single_numeric_values_are_unchanged(parse, value, expected):
    assert parse(value) == expected


@pytest.mark.parametrize('typed', [False, True])
def test_numeric_warnings_survive_model_normalization_and_serialization(typed):
    raw = {'product_name': '光跳纤架；光熔纤柜', 'quantity': '4套；5套',
           'unit_price': '319300；60480', 'total_price': '1277200；302400'}
    item = ItemExtract(**raw) if typed else raw
    result = ProjectExtract(lots=[{'lot_name': '合同包3', 'items': [item]}])
    saved = ProjectExtract.model_validate_json(result.model_dump_json())
    assert saved.processing_status == 'partial'
    assert len(saved.warnings) == 3
    assert all('分号并列数值' in warning for warning in saved.warnings)
    assert saved.lots[0].items[0].quantity is None
    assert saved.lots[0].items[0].unit_price is None
    assert saved.lots[0].items[0].total_price is None
    assert '_normalization_warnings' not in ItemExtract.model_json_schema()['properties']


def test_item_locations_survive_grounding_merging_scoring_and_storage(tmp_path):
    text = parse_file(table_file(tmp_path, [HEADER, GOODS]))
    result = ProjectExtract(notice_uuid='parallel-goods', lots=[LotExtract(lot_name='合同包3', items=[
        ItemExtract(product_name=name, model=model, quantity=quantity, unit_price=price,
                    supplier=GOODS[1], source_file='notice.html', line_id='TABLE:1 ROW:2')
        for name, model, quantity, price in [('光跳纤架', 'NGFC FRAME', 4, 319300), ('光熔纤柜', 'MX6', 5, 60480)]])])
    ground_evidence(result, text)
    result = merge_results([result, result.model_copy(deep=True)])
    result.notice_uuid = 'parallel-goods'
    assert [item.line_id for item in result.lots[0].items] == ['TABLE:1 ROW:2 ITEM:1', 'TABLE:1 ROW:2 ITEM:2']
    assert all(_verified(item) for item in result.lots[0].items)
    expected = [{**item.model_dump(), 'lot_name': '合同包3'} for item in result.lots[0].items]
    assert score_items(expected, expected[::-1])['micro']['f1'] == 1
    ids = upsert_notice(result, ['notice.html'])
    stored = rows('SELECT product_name,quantity,unit_price,total_price,source_location FROM items WHERE project_id=? ORDER BY id', (ids[0],))
    assert [(item['product_name'], item['quantity'], item['unit_price'], item['total_price']) for item in stored] == [
        ('光跳纤架', 4, 319300, None), ('光熔纤柜', 5, 60480, None)]
    assert [item['source_location'] for item in stored] == ['TABLE:1 ROW:2 ITEM:1', 'TABLE:1 ROW:2 ITEM:2']


def test_parser_ambiguity_is_reported_even_if_model_omits_warning(tmp_path, monkeypatch):
    data = [*GOODS]
    data[3] = '康普'
    path = table_file(tmp_path, [HEADER, data])
    extractor = DeepSeekExtractor()
    extractor.api_key = 'test-only'
    async def pretend_model(*args, **kwargs):
        return [ProjectExtract(lots=[LotExtract(lot_name='合同包3', items=[ItemExtract(product_name='设备')])])]
    monkeypatch.setattr(extractor, '_extract_piece', pretend_model)
    result = asyncio.run(extractor.extract([path]))
    assert result.processing_status == 'partial'
    assert any('保留原行未拆分' in warning for warning in result.warnings)


@pytest.mark.parametrize('current_cache', [False, True])
@pytest.mark.parametrize('cache_only', [False, True])
@pytest.mark.parametrize('text', ['[TABLE:1 ROW:2] 设备 | M1 | 1台 | 100元',
                                  '[TABLE:1 ROW:2 ITEM:1] 设备 | M1 | 1台 | 100元'])
def test_only_current_rules_can_match_model_cache(tmp_path, text, cache_only, current_cache):
    extractor = DeepSeekExtractor()
    extractor.settings = replace(extractor.settings, model_cache_dir=str(tmp_path))
    extractor.api_key = 'test-only'
    extractor.cache_only = cache_only
    payload = {'model': extractor.settings.deepseek_model,
               'messages': [{'role': 'system', 'content': system_prompt('schema') if current_cache else system_prompt('schema').removesuffix(MULTI_ITEM_RULES)},
                            {'role': 'user', 'content': user_prompt(text, ['notice.html'])}],
               'response_format': {'type': 'json_object'}, 'temperature': .1,
               'max_tokens': extractor.settings.model_max_tokens, 'stream': False}
    if extractor.settings.deepseek_model.startswith('deepseek-'):
        payload['thinking'] = {'type': 'enabled' if extractor.settings.model_thinking else 'disabled'}
    cache_key = hashlib.sha256((extractor.settings.deepseek_base_url + json.dumps(payload, sort_keys=True, ensure_ascii=False)).encode()).hexdigest()
    envelope = {'choices': [{'message': {'content': '{"project_name":"历史样本"}'}}]}
    (tmp_path / (cache_key + '.json')).write_text(json.dumps(envelope), encoding='utf-8')
    def handler(request):
        return httpx.Response(200, json={'choices': [{'message': {'content': '{"project_name":"新版抽取"}'}}]})
    async def run():
        extractor.client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
        try:
            return await extractor._call(text, ['notice.html'])
        finally:
            await extractor.close()
    if cache_only and not current_cache:
        with pytest.raises(FatalModelError, match='未命中当前规则'):
            asyncio.run(run())
    else:
        result = asyncio.run(run())
        assert result.project_name == ('历史样本' if current_cache else '新版抽取')
        assert not result.warnings


def test_pdf_parallel_goods_sharing_a_source_line_are_not_collapsed():
    text = '[PAGE:1 LINE:2] 光跳纤架；光熔纤柜 | NGFC FRAME；MX6 | 4套；5套 | 319300；60480'
    pieces = []
    for name, model in [('光跳纤架', 'NGFC FRAME'), ('光熔纤柜', 'MX6')]:
        piece = ProjectExtract(lots=[LotExtract(lot_name='合同包3', items=[
            ItemExtract(product_name=name, model=model, source_file='notice.pdf', line_id='PAGE:1 LINE:2')])])
        ground_evidence(piece, text)
        pieces.append(piece)
    result = merge_results([*pieces, *pieces])
    assert [item.product_name for item in result.lots[0].items] == ['光跳纤架', '光熔纤柜']
    assert result.processing_status == 'partial'
    assert any('保留独立记录' in warning for warning in result.warnings)


@pytest.mark.parametrize('difference', [
    {'brand': '品牌乙'}, {'quantity': 2}, {'unit_price': 200}, {'total_price': 200},
    {'quantity': 0},
])
def test_same_identity_pdf_goods_with_conflicting_fields_survive_storage_and_export(difference):
    first = ItemExtract(product_name='机柜', model='定制', brand='品牌甲', quantity=1,
                        unit_price=100, total_price=100, supplier='供应商甲',
                        source_file='notice.pdf', line_id='PAGE:1 LINE:2')
    second = first.model_copy(update=difference)
    pieces = [ProjectExtract(lots=[LotExtract(lot_name='合同包1', items=[item])])
              for item in (first, second)]
    # Repeated model chunks may repeat both goods; retain each distinct good once.
    result = merge_results([*pieces, *pieces])
    result.notice_uuid = 'same-identity-parallel-goods'
    assert len(result.lots[0].items) == 2
    assert result.processing_status == 'partial'
    assert any('保留独立记录' in warning for warning in result.warnings)
    upsert_notice(result, ['notice.pdf'])
    assert len(rows('SELECT id FROM items')) == 2
    records, notes = export_records('default')
    workbook = load_workbook(io.BytesIO(workbook_bytes(records, notes)))
    assert workbook['标的表'].max_row == 3
    assert records[0].processing_status == 'partial'


@pytest.mark.parametrize('parse,value,expected', [
    (money_to_yuan, '100万元；交付期90天', 1000000),
    (money_to_yuan, '交付期90天；100万元', 1000000),
    (money_to_yuan, '100元；90天内交货', 100),
    (money_to_yuan, '100元；质保3年', 100),
    (number_value, '4套；保修期：3年', 4),
    (number_value, '保修期：3年；4套', 4),
    (number_value, '服务期3年', 3),
])
def test_numeric_duration_notes_do_not_replace_or_clear_disclosed_values(parse, value, expected):
    assert parse(value) == expected


@pytest.mark.parametrize('raw', ['100元；其他费用200元', '100元；型号M2', '100元；200美元'])
def test_unknown_numeric_notes_remain_ambiguous(raw):
    item = ItemExtract(product_name='机柜', unit_price=raw)
    result = ProjectExtract(items=[item])
    assert item.unit_price is None
    assert result.processing_status == 'partial'
    assert len(result.warnings) == 1


@pytest.mark.parametrize('raw', ['100万元；200万元', '100万元；交付期90天；200万元'])
@pytest.mark.parametrize('typed', [False, True])
def test_parallel_project_lot_and_bid_amounts_are_reported_and_persisted(raw, typed):
    bid = {'bidder': '供应商甲', 'bid_amount': raw, 'score': '90；80'}
    lot = {'lot_name': '合同包1', 'budget_amount': raw, 'awarded_amount': raw,
           'items': [{'product_name': '机柜'}], 'bids': [BidExtract(**bid) if typed else bid]}
    result = ProjectExtract(notice_uuid='parallel-amounts', budget_amount=raw, awarded_amount=raw,
                            lots=[LotExtract(**lot) if typed else lot])
    assert result.processing_status == 'partial'
    assert len(result.warnings) == 6
    assert all('分号并列数值' in warning for warning in result.warnings)
    validate_result(result)
    upsert_notice(result, ['notice.html'])
    records, _ = export_records('default')
    saved = records[0]
    assert saved.processing_status == 'partial'
    assert saved.warnings == result.warnings
    assert saved.budget_amount is saved.awarded_amount is None
    assert saved.lots[0].budget_amount is saved.lots[0].awarded_amount is None
    assert saved.lots[0].bids[0].bid_amount is saved.lots[0].bids[0].score is None


def test_flat_bid_numeric_warnings_are_not_lost():
    result = ProjectExtract(bids=[BidExtract(bidder='供应商甲', bid_amount='100；200')])
    assert result.processing_status == 'partial'
    assert len(result.warnings) == 1


@pytest.mark.parametrize('ambiguous', [False, True])
def test_parse_only_reports_parser_warnings_to_job_and_document(tmp_path, ambiguous):
    data = [*GOODS]
    if ambiguous:
        data[3] = '康普'
    path = table_file(tmp_path, [HEADER, data])
    create_job('parallel-preflight', [path], tmp_path / 'job', groups=[[path]], mode='parse-only')
    asyncio.run(process_job('parallel-preflight'))
    state = get_job('parallel-preflight')
    document = rows('SELECT status,errors FROM job_documents')[0]
    assert state.metrics['model_calls'] == 0
    if ambiguous:
        assert state.status == 'completed_with_warnings' and state.partial == 1
        assert state.succeeded == 0 and state.failed == 0
        assert any('保留原行未拆分' in error for error in state.errors)
        assert document['status'] == 'partial'
        assert '保留原行未拆分' in document['errors']
    else:
        assert state.status == 'completed' and state.succeeded == 1
        assert state.partial == 0 and not state.errors
        assert document['status'] == 'parsed' and json.loads(document['errors']) == []


@pytest.mark.parametrize('html,expected', [
    ('<span>Model</span> <b>A</b>', 'Model A'),
    ('<span>Model </span><b>A</b>', 'Model A'),
    ('<span>Model</span>&nbsp;<b>A</b>', 'Model A'),
    ('<span>A</span><span>B</span>', 'AB'),
    ('<div>Model A</div><div>Model B</div>', 'Model A Model B'),
    ('<p>Model A</p><span>Model B</span>', 'Model A Model B'),
    ('Model A<br>Model B', 'Model A Model B'),
    ('<span>DH-ITSE0400-TA</span><br> <span>-N02</span>', 'DH-ITSE0400-TA-N02'),
    ('<span>DH-ITSE0400-TA-</span><br> <span>N02</span>', 'DH-ITSE0400-TA-N02'),
])
def test_html_keeps_source_word_boundaries_and_hyphen_continuations(tmp_path, html, expected):
    path = tmp_path / 'notice.html'
    path.write_text(f'<table><tr><td>{html}</td></tr></table>', encoding='utf-8')
    assert parse_file(path) == '[TABLE:1 ROW:1] ' + expected


def test_model_key_follows_merged_model_through_storage_and_search(client):
    def piece(model):
        return ProjectExtract(lots=[LotExtract(lot_name='合同包1', items=[
            ItemExtract(product_name='终端管理器', model=model, supplier='供应商甲',
                        source_file='notice.html', line_id='TABLE:1 ROW:2')])])
    result = merge_results([piece(''), piece('DH-ITSE0400-TA -N02')])
    result.notice_uuid = 'merged-model-key'
    assert result.lots[0].items[0].model_key == 'DH-ITSE0400-TA-N02'
    upsert_notice(result, ['notice.html'])
    found = client.get('/api/items', params={'model': 'DH-ITSE0400-TA-N02'}).json()
    assert found['total'] == 1
    assert found['data'][0]['model'] == 'DH-ITSE0400-TA -N02'


def test_model_key_is_recomputed_before_storing_an_existing_lot(client):
    item = ItemExtract(product_name='终端管理器', model='old-model')
    result = ProjectExtract(notice_uuid='changed-model-key', lots=[LotExtract(lot_name='包1', items=[item])])
    item.model = 'DH-ITSE0400-TA -N02'
    upsert_notice(result, ['notice.html'])
    assert rows('SELECT model_key FROM items')[0]['model_key'] == 'DH-ITSE0400-TA-N02'
    assert client.get('/api/items', params={'model': 'DH-ITSE0400-TA-N02'}).json()['total'] == 1
