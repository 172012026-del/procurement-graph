import io
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from app import analytics, database
from app.config import settings
from app.database import connect, init_db, one, rows, upsert_notice
from app.exports import workbook_bytes
from app.extractor import merge_results
from app.models import BidExtract, ItemExtract, LotExtract, ProjectExtract
from app.normalization import lot_identity


@pytest.mark.parametrize('name', ['包1', '合同包1', '第1包', '第一包', '采购包01', '合同包（1）'])
def test_explicit_number_aliases(name):
    assert lot_identity(name) == 'number:1'


def test_description_and_different_packages_not_collapsed():
    assert lot_identity('服务器包1') != lot_identity('网络包1')
    assert lot_identity('合同包1（设备采购）') == lot_identity('包1')
    assert lot_identity('合同包1（设备采购（一期））') == lot_identity('包1')
    assert lot_identity('A（设备采购）') == lot_identity('标包A')
    assert lot_identity('包1') != lot_identity('包2')
    assert lot_identity('包一二') == 'number:12'
    assert lot_identity('') == ''


def test_aliases_share_one_package_and_keep_source_label():
    records = [ProjectExtract(lot_name=name, items=[ItemExtract(product_name='服务器',
               source_file='notice.html', line_id='TABLE:1 ROW:2', brand=brand)])
               for name, brand in [('包1', ''), ('合同包1', '品牌甲')]]
    result = merge_results(records)
    assert len(result.lots) == 1 and len(result.lots[0].items) == 1
    assert result.lots[0].lot_name == '包1'
    assert result.lots[0].items[0].brand == '品牌甲'


def test_unknown_chunks_remain_separate_without_inventing_labels():
    result = merge_results([ProjectExtract(bids=[BidExtract(bidder=name, is_winner=True)]) for name in ('甲', '乙')])
    result.notice_uuid = 'unknown-notice'
    assert len(result.lots) == 2 and result.processing_status == 'partial'
    assert all(lot.lot_name == '' for lot in result.lots)
    upsert_notice(result, ['unknown-notice.html'])
    assert one('SELECT COUNT(*) n FROM projects')['n'] == 2
    assert analytics.common_projects(['甲', '乙'])['summary']['project_count'] == 0
    workbook = load_workbook(io.BytesIO(workbook_bytes([result])))
    assert workbook['总表'].max_row == 3
    assert workbook['总表']['D2'].value is None and workbook['总表']['D3'].value is None


def test_package_id_survives_alias_on_rerun(notice):
    first = upsert_notice(notice(lot='包1'), ['notice.html'])
    second = upsert_notice(notice(lot='合同包1'), ['notice.html'])
    assert first == second
    assert one('SELECT COUNT(*) n FROM projects')['n'] == 1
    assert one('SELECT lot_name FROM projects')['lot_name'] == '合同包1'


def test_v2_migration_preserves_alias_collisions_and_backups(tmp_path, monkeypatch):
    target = tmp_path / 'v2.db'
    monkeypatch.setattr(database, 'settings', replace(settings, database_path=target))
    init_db()
    with connect() as c:
        c.execute('DROP INDEX idx_notice_lot_key')
        c.execute('ALTER TABLE projects DROP COLUMN lot_key')
        c.execute("CREATE UNIQUE INDEX idx_notice_lot ON projects(dataset_id,notice_uuid,lot_name) WHERE notice_uuid<>''")
        for name in ('包1', '合同包1'):
            c.execute("INSERT INTO projects(name,dataset_id,notice_uuid,lot_name) VALUES('旧项目','official','n',?)", (name,))
    init_db()
    assert target.with_suffix('.pre-v3.db').is_file()
    stored = rows('SELECT lot_key,processing_status FROM projects ORDER BY id')
    assert len(stored) == 2 and len({r['lot_key'] for r in stored}) == 2
    assert stored[1]['processing_status'] == 'partial'
