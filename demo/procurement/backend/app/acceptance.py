"""Local acceptance checks backed by the organizer's sheet schema and source facts."""
from __future__ import annotations

from collections import Counter, defaultdict
from decimal import Decimal
import hashlib
import io
from itertools import combinations
import json
import math
from pathlib import Path

from openpyxl import load_workbook

from . import analytics
from .database import rows
from .normalization import lot_identity, normalize_org


def check_workbook(content: bytes, schema_path: Path, notice_ids: set[str] | None = None) -> dict:
    schema = json.loads(schema_path.read_text('utf-8-sig'))
    workbook = load_workbook(io.BytesIO(content), data_only=False, read_only=True)
    errors, blanks, counts = [], {}, {}
    expected_names = [sheet['sheet'] for sheet in schema]
    if workbook.sheetnames != expected_names:
        errors.append('Sheet names/order differ from the official three-sheet template')
    package_keys = set()
    for spec in schema:
        name = spec['sheet']
        if name not in workbook:
            continue
        sheet = workbook[name]
        fields = spec['fields']
        headers = [cell.value for cell in next(sheet.iter_rows())]
        if headers != [field['name'] for field in fields]:
            errors.append(f'{name}: headers differ from the official template')
        counts[name] = max(0,sheet.max_row-1)
        blanks[name] = {field['name']:0 for field in fields if field['required_blue']}
        uuid_column = next(i for i,f in enumerate(fields) if f['name']=='UUID')
        lot_column = next(i for i,f in enumerate(fields) if f['name']=='标段/合同包')
        for row_number,cells in enumerate(sheet.iter_rows(min_row=2),2):
            values = [cell.value for cell in cells]
            uid, label = values[uuid_column], values[lot_column] or ''
            if notice_ids is not None and uid not in notice_ids:
                errors.append(f'{name}:{row_number}: UUID not present in the source input')
            key = (uid,label)
            if name == '总表':
                if key in package_keys:
                    errors.append(f'{name}:{row_number}: duplicate announcement/package row')
                package_keys.add(key)
            elif key not in package_keys:
                errors.append(f'{name}:{row_number}: no matching package in the summary sheet')
            for index,(cell,field) in enumerate(zip(cells,fields)):
                if cell.data_type == 'f':
                    errors.append(f'{name}:{row_number}: formula cell is not a source fact')
                if field['required_blue'] and cell.value in (None,''):
                    blanks[name][field['name']] += 1
                numeric = field['name'] in {'中标金额','单价','总价','数量','报价（万元）'}
                if numeric and cell.value is not None:
                    if not isinstance(cell.value,(int,float)) or not math.isfinite(cell.value):
                        errors.append(f'{name}:{row_number}:{field["name"]}: invalid numeric value')
                    elif field['name'] in {'中标金额','报价（万元）'}:
                        value = Decimal(str(cell.value))
                        if value != value.quantize(Decimal('.01')):
                            errors.append(f'{name}:{row_number}:{field["name"]}: expected two decimal places')
    workbook.close()
    return {'valid':not errors,'errors':errors,'row_counts':counts,'required_blue_positions':sum(sum(f['required_blue'] for f in s['fields']) for s in schema),
        'blank_cells':blanks,'blank_policy':'Source omissions remain blank; a blue column does not authorize inventing a value.',
        'scope':'Three-sheet structure, types, UUID/package links and rounding; semantic accuracy requires source labels.',
        'schema_sha256':hashlib.sha256(schema_path.read_bytes()).hexdigest(),'workbook_sha256':hashlib.sha256(content).hexdigest()}


def check_relationships(reference_path: Path, dataset_id: str) -> dict:
    source = [json.loads(line) for line in reference_path.read_text('utf-8-sig').splitlines() if line.strip()]
    for record in source:
        for package in record['packages']:
            package['winners'] = [normalize_org(name) for name in package['winners']]
            for bid in package['bids']:
                bid['bidder'] = normalize_org(bid['bidder'])
    packages = {(r['notice_uuid'],lot_identity(p['lot_name'])):{**p,'purchaser':normalize_org(r['purchaser'])}
        for r in source for p in r['packages']}
    actual_packages = {(p['notice_uuid'],p['lot_key']):p for p in rows('SELECT * FROM projects WHERE dataset_id=?',(dataset_id,))}
    errors = []
    for key,p in packages.items():
        actual = actual_packages.get(key)
        if not actual:
            errors.append({'package':key,'field':'package','expected':'present','actual':'missing'})
            continue
        for field in ('purchaser','awarded_amount'):
            if actual[field] != p[field]:
                errors.append({'package':key,'field':field,'expected':p[field],'actual':actual[field]})
        bids = rows('SELECT * FROM bids WHERE project_id=?',(actual['id'],))
        for field,predicate in [('winners',lambda b:b['is_winner']),('eligible_bidders',lambda b:b['eligible'])]:
            expected = set(p['winners']) if field=='winners' else {b['bidder'] for b in p['bids'] if b['eligible']}
            found = {b['bidder'] for b in bids if predicate(b)}
            if expected != found:
                errors.append({'package':key,'field':field,'expected':sorted(expected),'actual':sorted(found)})

    buyers = sorted({p['purchaser'] for p in packages.values()})
    buyer = max(buyers,key=lambda b:len({w for p in packages.values() if p['purchaser']==b for w in p['winners']}))
    selected = {k:p for k,p in packages.items() if p['purchaser']==buyer}
    winner_names = sorted({w for p in selected.values() for w in p['winners']})
    all_winners = {w for p in packages.values() for w in p['winners']}
    eligible = {key:{b['bidder'] for b in p['bids'] if b['eligible']} for key,p in packages.items()}
    cooccurrences = Counter(pair for names in eligible.values() for pair in combinations(sorted(names & all_winners),2))
    common_pair = list(max(cooccurrences,key=cooccurrences.get)) if cooccurrences else winner_names[:2]
    competitor_supplier = max(all_winners,key=lambda s:sum(s in names for names in eligible.values()))
    cases = []

    def record(name,params,expected,actual):
        cases.append({'scenario':name,'parameters':params,'expected':expected,'actual':actual,'passed':expected==actual})

    partners = analytics.purchaser_partners(buyer,dataset_id)
    expected = {name:{'count':sum(name in p['winners'] for p in selected.values()),
        'amount':sum(p['awarded_amount'] or 0 for p in selected.values() if name in p['winners'])} for name in winner_names}
    actual = {r['supplier']:{'count':r['cooperation_count'],'amount':r['total_amount'] or 0} for r in partners['winning_suppliers']}
    record('purchaser-partners',{'purchaser':buyer},expected,actual)

    frequencies = Counter(name for key in selected for name in eligible[key])
    data = analytics.purchaser_bidders(buyer,dataset_id)
    top = data['top_bidders']
    cutoff = sorted(frequencies.values(),reverse=True)[min(4,len(frequencies)-1)] if frequencies else 0
    pairs = Counter(pair for key in selected for pair in combinations(sorted(eligible[key]),2))
    top_valid = len(top)==min(5,len(frequencies)) and all(frequencies[r['bidder']]==r['frequency'] and r['frequency']>=cutoff for r in top)
    pair_valid = all(pairs[(r['bidder_a'],r['bidder_b'])]==r['frequency'] for r in data['collaborative_pairs'])
    pair_valid &= len(data['collaborative_pairs'])==min(10,len(pairs))
    record('purchaser-bidders',{'purchaser':buyer},{'top5_counts_and_pairs_valid':True}, {'top5_counts_and_pairs_valid':bool(top_valid and pair_valid)})

    own = {key for key,names in eligible.items() if competitor_supplier in names}
    counts = Counter(name for key in own for name in eligible[key] if name != competitor_supplier)
    data = analytics.supplier_competitors(competitor_supplier,dataset_id)
    record('supplier-competitors',{'supplier':competitor_supplier},dict(counts),{r['bidder']:r['frequency'] for r in data['all_participants']})

    pair = winner_names[:2]
    expected = {}
    for purchaser in buyers:
        found = {key:p for key,p in packages.items() if p['purchaser']==purchaser and set(p['winners']) & set(pair)}
        if set(pair).issubset({w for p in found.values() for w in p['winners']}):
            expected[purchaser] = {'count':len(found),'amount':round(sum(p['awarded_amount'] or 0 for p in found.values()),2)}
    data = analytics.common_purchasers(pair,dataset_id)
    record('common-purchasers',{'suppliers':pair},expected,{r['purchaser']:{'count':r['cooperation_count'],'amount':r['total_amount'] or 0} for r in data['purchasers']})

    found = {key:p for key,p in packages.items() if set(common_pair).issubset(eligible[key])}
    data = analytics.common_projects(common_pair,dataset_id)
    expected = {'packages':sorted('|'.join(key) for key in found),'count':len(found),'amount':round(sum(p['awarded_amount'] or 0 for p in found.values()),2)}
    actual = {'packages':sorted(r['notice_uuid']+'|'+r['lot_key'] for r in data['projects']),
        'count':data['summary']['project_count'],'amount':data['summary']['total_amount'] or 0}
    record('common-projects',{'suppliers':common_pair},expected,actual)
    return {'passed':not errors and all(c['passed'] for c in cases),'dataset_id':dataset_id,'official_score':False,
        'identity_policy':'NFKC and whitespace normalization; source labels remain in the reference file.',
        'source_notices':len(source),'source_packages':len(packages),'source_bid_rows':sum(len(p['bids']) for p in packages.values()),
        'reference_sha256':hashlib.sha256(reference_path.read_bytes()).hexdigest(),'fact_errors':errors,'scenarios':cases,
        'scope':'Local source-table checks for five business scenarios; not the organizer hidden-test evaluator.'}
