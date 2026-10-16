"""Organize supplied competition inputs and inventory them without model calls.

Run with Python 3.12 from the project root. Source archives stay in the source
directory. Only HTMLs, notice-level attachment ZIPs, and the small format
references are copied into data/official-benchmark. Attachment ZIPs are not expanded.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sys
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path, PurePosixPath

from lxml import html as lxml_html
from openpyxl import load_workbook


WORKSPACE = Path(__file__).resolve().parents[1]
SOURCE = Path.home() / 'Downloads'
DESTINATION = WORKSPACE / "data/official-benchmark"
INPUTS = [
    "赛题五.基准测试数据_html.zip",
    "赛题五.基准测试数据_file.zip",
    "关于任务数据提取结果格式的要求及补充说明.pdf",
    "输出标准样例表.xlsx",
]
REFERENCE_NAMES = {
    INPUTS[2]: 'output-format.pdf',
    INPUTS[3]: 'output-template.xlsx',
}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(4 * 1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_json(path: Path, data: object) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(path: Path, data: list[dict], columns: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(data)


def relative(path: Path) -> str:
    return path.relative_to(WORKSPACE).as_posix()


def extract_layer(source: Path, destination: Path, suffix: str) -> list[dict]:
    records: list[dict] = []
    destination.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source) as archive:
        entries = [x for x in archive.infolist() if not x.is_dir()]
        for index, entry in enumerate(entries, 1):
            name = PurePosixPath(entry.filename.replace("\\", "/"))
            if name.is_absolute() or ".." in name.parts or ":" in name.parts[0]:
                raise ValueError(f"Unsafe archive member: {entry.filename}")
            if name.suffix.lower() != suffix:
                raise ValueError(f"Unexpected archive member: {entry.filename}")
            target = destination.joinpath(*name.parts)
            if not target.resolve().is_relative_to(destination.resolve()):
                raise ValueError(f"Archive member escaped destination: {entry.filename}")
            target.parent.mkdir(parents=True, exist_ok=True)
            h = hashlib.sha256()
            if target.exists():
                if target.stat().st_size != entry.file_size:
                    raise ValueError(f"Existing destination differs: {target}")
                with archive.open(entry) as original, target.open("rb") as existing:
                    while chunk := original.read(1024 * 1024):
                        if existing.read(len(chunk)) != chunk:
                            raise ValueError(f"Existing destination differs: {target}")
                        h.update(chunk)
            else:
                partial = target.with_name(target.name + ".part")
                with archive.open(entry) as original, partial.open("wb") as output:
                    while chunk := original.read(1024 * 1024):
                        output.write(chunk)
                        h.update(chunk)
                partial.replace(target)
            records.append({"uuid": target.stem, "path": relative(target),
                            "bytes": entry.file_size, "sha256": h.hexdigest(),
                            "outer_member": entry.filename})
            if index % 100 == 0 or index == len(entries):
                print(f"{source.name}: {index}/{len(entries)} files", flush=True)
    return records


def display_name(name: str) -> str:
    try:
        raw = name.encode("cp437")
    except UnicodeEncodeError:
        return name
    for encoding in ("utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    return name


def workbook_schema(path: Path) -> list[dict]:
    book = load_workbook(path, data_only=False)
    sheets = []
    for sheet in book:
        fields = []
        for cell in sheet[1]:
            if cell.value is None:
                continue
            color = cell.fill.fgColor
            blue = color.type == "rgb" and color.rgb == "FF1F4E79"
            fields.append({"column": cell.column_letter, "name": cell.value,
                           "required_blue": blue,
                           "instruction_cell": sheet.cell(2, cell.column).coordinate,
                           "instruction": sheet.cell(2, cell.column).value,
                           "example": sheet.cell(3, cell.column).value})
        samples = [list(row) for row in sheet.iter_rows(min_row=4, values_only=True)
                   if any(value is not None for value in row)]
        sheets.append({"sheet": sheet.title, "used_range": sheet.calculate_dimension(),
                       "required_blue_count": sum(x["required_blue"] for x in fields),
                       "fields": fields, "sample_row_count": len(samples),
                       "sample_rows_first_two": samples[:2]})
    book.close()
    return sheets


def inspect_attachment_headers(root: Path) -> dict:
    """Read at most 4096 bytes/member; do not expand attachment payloads."""
    index_dir = root / "indexes"
    with (index_dir / "attachment-archives.csv").open(encoding="utf-8-sig") as stream:
        packages = list(csv.DictReader(stream))
    records = []
    expected = {".pdf": {"pdf"}, ".zip": {"zip"}, ".rar": {"rar"},
                ".7z": {"7z"}, ".png": {"png"}, ".jpg": {"jpeg"},
                ".jpeg": {"jpeg"}, ".docx": {"zip"}, ".xlsx": {"zip"},
                ".doc": {"ole"}, ".xls": {"ole"}, ".wps": {"ole", "zip"}}
    for index, package in enumerate(packages, 1):
        if package["zip_state"] != "zip_valid":
            continue
        with zipfile.ZipFile(WORKSPACE / package["path"]) as archive:
            for entry in archive.infolist():
                if entry.is_dir():
                    continue
                name = display_name(entry.filename)
                ext = PurePosixPath(name).suffix.lower() or "[无扩展名]"
                error = ""
                try:
                    with archive.open(entry) as stream:
                        head = stream.read(4096)
                    trimmed = head.lstrip(b"\xef\xbb\xbf\x00 \r\n\t")
                    if head.startswith(b"%PDF-"):
                        kind = "pdf"
                    elif head.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
                        kind = "zip"
                    elif head.startswith(b"Rar!\x1a\x07"):
                        kind = "rar"
                    elif head.startswith(b"7z\xbc\xaf\x27\x1c"):
                        kind = "7z"
                    elif head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
                        kind = "ole"
                    elif head.startswith(b"\x89PNG\r\n\x1a\n"):
                        kind = "png"
                    elif head.startswith(b"\xff\xd8\xff"):
                        kind = "jpeg"
                    elif re.search(br"<!doctype\s+html|<(?:html|head|body|title|link|script|div|meta)(?:\s|>)", trimmed, re.I):
                        kind = "html"
                    elif not head:
                        kind = "empty"
                    else:
                        kind = "other_or_unknown"
                except Exception as exc:
                    kind = "read_error"
                    error = type(exc).__name__ + ": " + str(exc)[:160]
                mismatch = ext in expected and kind not in expected[ext]
                records.append({"uuid": package["uuid"], "member": name,
                                "extension": ext, "bytes": entry.file_size,
                                "detected_header": kind, "extension_mismatch": mismatch,
                                "error": error})
        if index % 200 == 0 or index == len(packages):
            print(f"Attachment file headers inspected: {index}/{len(packages)} packages", flush=True)
    columns = ["uuid", "member", "extension", "bytes", "detected_header", "extension_mismatch", "error"]
    write_csv(index_dir / "attachment-signatures.csv", records, columns)
    flagged = [x for x in records if x["extension_mismatch"] or x["detected_header"] in {"html", "empty", "read_error"}]
    write_csv(index_dir / "attachment-content-anomalies.csv", flagged, columns)
    result = {"members_inspected": len(records),
              "detected_headers": dict(Counter(x["detected_header"] for x in records)),
              "flagged_members": len(flagged),
              "extension_mismatches": sum(x["extension_mismatch"] for x in records),
              "scope": "First 4096 bytes per direct member, no recursive archive or full-file validation"}
    summary_path = index_dir / "data-health.json"
    summary = json.loads(summary_path.read_text("utf-8"))
    summary["attachment_header_review"] = result
    write_json(summary_path, summary)
    print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
    return result


def html_profile(path: Path) -> dict:
    document = lxml_html.fromstring(path.read_bytes())
    title = " ".join(document.xpath("//title/text()"))
    content = document.xpath("//*[@class='vF_detail_content' or @id='info_content']")
    content = content[0] if content else document
    text = " ".join(content.itertext())
    urls = [str(x) for x in document.xpath("//a/@href")]
    file_links = [u for u in urls if re.search(r"\.(zip|rar|pdf|docx?|xlsx?|jpg|png)(?:$|[?#])", u, re.I)]
    kind = "其他或待确认"
    for pattern, category in [(r"更正|变更", "更正/变更"), (r"终止|废标|流标", "终止/废标"),
                              (r"中标|成交|结果", "中标/成交/结果"),
                              (r"招标|磋商|谈判|询价|采购公告", "招标/采购")]:
        if re.search(pattern, title):
            kind = category
            break
    canonical = document.xpath("//link[@rel='canonical']/@href")
    return {"title": title, "title_category": kind, "visible_characters": len(text),
            "table_count": len(content.xpath(".//table")), "file_link_count": len(file_links),
            "package_keyword": bool(re.search("合同包|采购包|标段", text)),
            "review_keyword": bool(re.search("资格性|符合性|资格审查|符合性审查", text)),
            "score_keyword": bool(re.search("综合得分|得分排名|评审得分|综合评分", text)),
            "canonical_url": canonical[0] if canonical else ""}


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    root = DESTINATION.resolve()
    if not root.is_relative_to(WORKSPACE.resolve()):
        raise ValueError("Destination must remain in workspace")
    if shutil.disk_usage(WORKSPACE).free < 6 * 1024**3:
        raise RuntimeError("Need at least 6 GiB free to organize these inputs")
    for name in INPUTS:
        if not (SOURCE / name).is_file():
            raise FileNotFoundError(SOURCE / name)
    index_dir = root / "indexes"
    refs = root / "format"
    index_dir.mkdir(parents=True, exist_ok=True)
    refs.mkdir(parents=True, exist_ok=True)
    for name in INPUTS[2:]:
        target = refs / REFERENCE_NAMES[name]
        if target.exists() and digest(target) != digest(SOURCE / name):
            raise ValueError(f"Existing reference differs: {target}")
        if not target.exists():
            shutil.copyfile(SOURCE / name, target)
    htmls = extract_layer(SOURCE / INPUTS[0], root / "html", ".html")
    attachments = extract_layer(SOURCE / INPUTS[1], root / "file", ".zip")
    attachment_by_uuid = {x["uuid"]: x for x in attachments}
    html_by_uuid = {x["uuid"]: x for x in htmls}
    members: list[dict] = []
    anomalies: list[dict] = []
    formats: Counter = Counter()
    package_states: Counter = Counter()
    for index, record in enumerate(attachments, 1):
        path = WORKSPACE / record["path"]
        try:
            with zipfile.ZipFile(path) as archive:
                files = [x for x in archive.infolist() if not x.is_dir()]
                state = "zip_valid" if files else "zip_empty"
                record.update(zip_state=state, member_count=len(files),
                              declared_expanded_bytes=sum(x.file_size for x in files))
                package_states[state] += 1
                for entry in files:
                    name = display_name(entry.filename)
                    ext = PurePosixPath(name).suffix.lower() or "[无扩展名]"
                    formats[ext] += 1
                    members.append({"uuid": record["uuid"], "package_path": record["path"],
                                    "member": name, "extension": ext, "bytes": entry.file_size,
                                    "compressed_bytes": entry.compress_size,
                                    "encrypted": bool(entry.flag_bits & 1)})
                if not files:
                    anomalies.append({"uuid": record["uuid"], "path": record["path"],
                                      "issue": "空ZIP", "detail": "中央目录无文件"})
        except (zipfile.BadZipFile, OSError, NotImplementedError) as exc:
            with path.open("rb") as stream:
                start = stream.read(300)
            text_like = bool(re.search(br"html|doctype|404|403|error", start, re.I))
            record.update(zip_state="invalid_zip", member_count=0, declared_expanded_bytes=0)
            package_states["invalid_zip"] += 1
            anomalies.append({"uuid": record["uuid"], "path": record["path"],
                              "issue": "不是可读取ZIP", "detail": type(exc).__name__ +
                              ("; 文件头疑似网页或错误响应" if text_like else "")})
        if index % 200 == 0 or index == len(attachments):
            print(f"Attachment central directories inspected: {index}/{len(attachments)}", flush=True)
    notices = []
    for record in htmls:
        attachment = attachment_by_uuid.get(record["uuid"])
        try:
            profile = html_profile(WORKSPACE / record["path"])
        except Exception as exc:
            profile = {"title": "", "title_category": "HTML解析异常", "error": type(exc).__name__}
        notices.append({**record, **profile,
                        "attachment_path": attachment["path"] if attachment else "",
                        "attachment_bytes": attachment["bytes"] if attachment else 0,
                        "attachment_state": attachment["zip_state"] if attachment else "无同名附件包",
                        "attachment_members": attachment["member_count"] if attachment else 0})
    sources = []
    for name in INPUTS:
        print(f"Hashing source: {name}", flush=True)
        path = SOURCE / name
        item = {"source": str(path), "bytes": path.stat().st_size, "sha256": digest(path)}
        if name in INPUTS[2:]:
            copied = refs / REFERENCE_NAMES[name]
            item.update(local_copy=relative(copied), copy_sha256=digest(copied))
        sources.append(item)
    schema = workbook_schema(refs / REFERENCE_NAMES[INPUTS[3]])
    summary = {
        "generated_at": datetime.now().astimezone().isoformat(),
        "html_count": len(htmls), "attachment_package_count": len(attachments),
        "paired_count": len(set(html_by_uuid) & set(attachment_by_uuid)),
        "html_without_package": sorted(set(html_by_uuid) - set(attachment_by_uuid)),
        "packages_without_html": sorted(set(attachment_by_uuid) - set(html_by_uuid)),
        "duplicate_html_stems": len(htmls) - len(html_by_uuid),
        "duplicate_package_stems": len(attachments) - len(attachment_by_uuid),
        "html_duplicate_content_count": len(htmls) - len({x['sha256'] for x in htmls}),
        "package_states": dict(package_states), "member_count": len(members),
        "member_extensions": dict(formats.most_common()),
        "declared_inner_expanded_bytes": sum(x["bytes"] for x in members),
        "max_attachment_package_bytes": max(x["bytes"] for x in attachments),
        "attachment_packages_above_200_mib": sum(x["bytes"] > 200 * 1024**2 for x in attachments),
        "title_categories": dict(Counter(x.get("title_category") for x in notices)),
        "html_with_package_keyword": sum(bool(x.get("package_keyword")) for x in notices),
        "html_with_review_keyword": sum(bool(x.get("review_keyword")) for x in notices),
        "html_with_score_keyword": sum(bool(x.get("score_keyword")) for x in notices),
        "html_without_package_but_with_file_links": [x["uuid"] for x in notices
            if not x["attachment_path"] and x.get("file_link_count", 0) > 0],
        "template_required_blue_counts": {x["sheet"]: x["required_blue_count"] for x in schema},
        "validation_scope": {"outer_member_crc": "verified while reading and extracting",
                             "inner_zip_validation": "central directory only; member data CRC not exhaustively checked",
                             "nested_archives": "inventoried as members; not recursively expanded",
                             "model_calls": 0, "semantic_extraction": False},
    }
    write_json(index_dir / "sources-and-checksums.json", sources)
    write_json(index_dir / "output-schema.json", schema)
    write_json(index_dir / "data-health.json", summary)
    write_csv(index_dir / "notice-attachment-manifest.csv", notices, ["uuid", "path", "bytes", "sha256",
              "title", "title_category", "attachment_path", "attachment_bytes", "attachment_state",
              "attachment_members", "visible_characters", "table_count", "file_link_count",
              "package_keyword", "review_keyword", "score_keyword", "canonical_url"])
    write_csv(index_dir / "attachment-archives.csv", attachments, ["uuid", "path", "bytes", "sha256",
              "zip_state", "member_count", "declared_expanded_bytes"])
    write_csv(index_dir / "attachment-members.csv", members, ["uuid", "package_path", "member",
              "extension", "bytes", "compressed_bytes", "encrypted"])
    write_csv(index_dir / "attachment-anomalies.csv", anomalies, ["uuid", "path", "issue", "detail"])
    print(json.dumps({k: v for k, v in summary.items() if not isinstance(v, list)}, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, default=SOURCE)
    parser.add_argument('--headers-only', action='store_true')
    args = parser.parse_args()
    SOURCE = args.source_dir.resolve()
    if args.headers_only:
        sys.stdout.reconfigure(encoding="utf-8")
        inspect_attachment_headers(DESTINATION)
    else:
        main()
        inspect_attachment_headers(DESTINATION)
